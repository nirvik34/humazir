
import json
import re
from datetime import datetime
from pathlib import Path

from humanizr.config import CFG
from humanizr.utils import ensure_dir


def save_report(data: dict, stem="report") -> str:
    out_dir = ensure_dir(CFG["output"]["reports_dir"])

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"{stem}_{ts}.json"

    path.write_text(
        json.dumps(data, indent=2),
        encoding="utf-8"
    )

    return str(path)


def build_detect_report(file_path, result):
    return {
        "tool": "humanizr",
        "command": "detect",
        "timestamp": datetime.now().isoformat(),
        "file": str(file_path),
        "ai_score": result.ai_score,
        "confidence": result.confidence,
        "verdict": result.verdict,
        "signals": result.signals,
        "human_signals": result.human_signals,
        "flagged_passages": result.flagged_passages
    }


def build_humanize_report(file_path, result):
    return {
        "tool": "humanizr",
        "command": "humanize",
        "timestamp": datetime.now().isoformat(),
        "file": str(file_path),
        "style": result.style,
        "word_count_before": result.word_count_before,
        "word_count_after": result.word_count_after,
        "changes": result.changes
    }


def export_highlighted_txt(original_text, flagged_passages, stem="file"):
    out_dir = ensure_dir(CFG["output"]["exports_dir"])

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"{stem}_marked_{ts}.txt"

    open_mark = CFG["output"]["highlight_marker_open"]
    close_mark = CFG["output"]["highlight_marker_close"]

    flagged = {p["line"] for p in flagged_passages}

    sents = [
        s.strip()
        for s in re.split(r'(?<=[.!?])\s+', original_text)
        if s.strip()
    ]

    output = []

    for i, sent in enumerate(sents, 1):
        if i in flagged:
            output.append(f"{open_mark}{sent}{close_mark}")
        else:
            output.append(sent)

    path.write_text(" ".join(output), encoding="utf-8")

    return str(path)


def export_humanized_txt(text, stem="file"):
    out_dir = ensure_dir(CFG["output"]["exports_dir"])

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"{stem}_humanized_{ts}.txt"

    path.write_text(text, encoding="utf-8")

    return str(path)


def export_highlighted_docx(original_text, flagged_passages, stem="file"):
    try:
        from docx import Document
        from docx.shared import RGBColor
    except ImportError:
        raise ImportError(
            "python-docx required.\nRun: pip install python-docx"
        )

    out_dir = ensure_dir(CFG["output"]["exports_dir"])

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = out_dir / f"{stem}_marked_{ts}.docx"

    flagged = {p["line"] for p in flagged_passages}

    sents = [
        s.strip()
        for s in re.split(r'(?<=[.!?])\s+', original_text)
        if s.strip()
    ]

    doc = Document()

    doc.add_heading("Humanizr Detection Report", level=1)
    doc.add_paragraph(
        f"Flagged Sentences: {len(flagged)}"
    )

    for i, sent in enumerate(sents, 1):
        para = doc.add_paragraph()
        run = para.add_run(sent)

        if i in flagged:
            run.bold = True
            run.font.color.rgb = RGBColor(200, 0, 0)

    doc.save(str(path))

    return str(path)